"""
catalog.py - description.xlsx checked against what is really in data/input, as a new Excel.

  uv run pipeline.py --mode all --vision    # first: transcripts (Whisper where missing), slides, search
  uv run catalog.py                         # then: data/output/description_checked.xlsx
  uv run catalog.py --no-analyse            # only what is there (no model calls, free)

Your description.xlsx is not changed. The copy gets these columns:
  In data/input   video / video + audio / audio only / transcript only / MISSING
  Files           the recording and transcript files found for the row
  Transcript      .vtt / .docx (given), made with Whisper (path), or NONE yet
  Slides          how many slides were read from the video (text), and described (--vision)
  Minutes         length of the recording (or of the transcript)
  Topics          what the recording covers, from its transcript and slides (topics.py)
  Check           anything to look at (wrong pair, other language, not processed yet ...)
and one new row (yellow) for every recording or transcript that is not in the Excel yet: the id
from the file name, the name, and a 1-3 sentence description written by the model from the
transcript and the slides. Check them; then the file can replace description.xlsx (extra
columns are ignored there).

The analysis is topics.py: cached per recording, so a second run costs nothing. Roughly 1-3
cents per 10 minutes of recording, once.
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import config
import inventory
import llm
import store

log = logging.getLogger("catalog")

COLUMNS = ["In data/input", "Files", "Transcript", "Slides", "Minutes", "Topics (from the recording)", "Check"]
ADDED = "ADDED: not in the Excel before; description written by AI from the transcript and slides - please check"
ORDER = {"video": 0, "audio": 1, "none": 2}
FILL = {"new": "FFF2CC", "missing": "F8CBAD", "ok": "E2EFDA"}


def recordings() -> list[dict]:
    """Manifest rows of recordings and of transcripts without a recording (not documents)."""
    return [r for r in store.read_manifest() if r.get("media_type") in ORDER]


def entry_of(row: dict, entries: dict) -> inventory.Entry | None:
    """The Excel row a recording belongs to (by the number in the file name, else by name)."""
    for name in (row.get("video_file"), row.get("transcript_file")):
        if name:
            entry, _ = inventory.match_entry(Path(name).name, entries)
            if entry:
                return entry
    return None


def presence(recs: list[dict]) -> str:
    kinds = {r.get("media_type") for r in recs}
    if "video" in kinds:
        return "video + audio" if "audio" in kinds else "video"
    if "audio" in kinds:
        return "audio only"
    return "transcript only" if recs else "MISSING"


def transcript_info(row: dict) -> str:
    """Where the transcript comes from: given (.vtt / .docx), made with Whisper, or none yet."""
    if row.get("transcript_file"):
        return f"{Path(row['transcript_file']).suffix.lower()} (given)"
    made = store.generated_transcript(row["video_id"])
    if made.exists():
        try:
            made = made.relative_to(config.ROOT)
        except ValueError:
            pass
        return f"made with Whisper: {Path(made).as_posix()}"
    return "NONE yet - run: uv run pipeline.py --mode all"


def slides_info(row: dict) -> str:
    vid = row["video_id"]
    if not store.has_slides(vid):
        return "not read yet (uv run pipeline.py)"
    slides = [s for s in store.load_slides(vid) if s.get("is_slide")]
    described = sum(bool(s.get("visual")) for s in slides)
    if not slides:
        return "no slides on screen"
    return f"{len(slides)} slides read" + (f", {described} described" if described else " (add --vision to describe them)")


def minutes(recs: list[dict]) -> float | str:
    values = []
    for r in recs:
        for key in ("video_min", "transcript_min"):
            try:
                values.append(float(r.get(key) or 0))
            except ValueError:
                pass
    return max(values, default=0) or ""


class Analyser:
    """topics.py per recording (summary + topics). Stops calling the model once the spending lock
    is reached; the Excel is still written."""

    def __init__(self, on: bool):
        self.on, self.stopped = on, ""

    def __call__(self, row: dict) -> dict:
        if not self.on:
            return {}
        if row.get("include") != "yes":
            return {"note": "include=no in the manifest: not analysed"}
        if not store.has_segments(row["video_id"]):
            return {"note": "not processed yet: run uv run pipeline.py --mode all, then catalog.py again"}
        if self.stopped:
            return {"note": f"not analysed: {self.stopped}"}
        import topics
        try:
            return topics.analyse_recording(row)
        except llm.BudgetExceeded as e:
            self.stopped = str(e)[:120]
            log.error("Analysis stopped: %s", e)
            return {"note": f"not analysed: {self.stopped}"}
        except Exception as e:   # one recording must not stop the others
            log.error("%s: analysis failed - %s", row.get("title"), e)
            return {"note": f"analysis failed: {str(e)[:120]}"}


def _header(wb) -> tuple:
    """The sheet and row with the 'id' and 'name' headers (like inventory.load_descriptions)."""
    for ws in wb.worksheets:
        for cells in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
            names = [" ".join(str(c.value).split()).lower() if c.value is not None else "" for c in cells]
            if "id" in names and "name" in names:
                cols = {k: names.index(k) + 1 for k in ("id", "name", "description") if k in names}
                last = max(i + 1 for i, n in enumerate(names) if n)
                return ws, cells[0].row, cols, last
    raise ValueError("no sheet has a header row with the columns 'id' and 'name'")


def build(analyse: bool = True) -> tuple[Path, dict]:
    """Write the checked copy of description.xlsx. Returns (path, counts)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = inventory.find_descriptions_file()
    if not path or not path.exists():
        raise FileNotFoundError(f"No description*.xlsx in {config.INPUT_DIR} (or DESCRIPTIONS_FILE in .env)")
    entries = inventory.load_descriptions(path)
    rows = recordings()
    if not rows:
        raise RuntimeError("The manifest has no recordings yet. First: uv run pipeline.py --mode all")
    by_id: dict[str, list[dict]] = {}
    added: list[dict] = []
    for r in rows:
        e = entry_of(r, entries)
        if e:
            by_id.setdefault(e.id, []).append(r)
        elif "AUDIO OF A VIDEO" not in (r.get("check") or ""):   # the audio copy of a video is not a new item
            added.append(r)

    wb = load_workbook(path)
    ws, head, cols, last = _header(wb)
    if "description" not in cols:
        last += 1
        cols["description"] = last
        ws.cell(head, last, "description").font = Font(bold=True)
    col = {name: last + 1 + i for i, name in enumerate(COLUMNS)}
    for name, c in col.items():
        ws.cell(head, c, name).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = 45 if name in ("Files", "Topics (from the recording)",
                                                                             "Check", "Transcript") else 18
    analyser = Analyser(analyse)
    wrap = Alignment(wrap_text=True, vertical="top")

    def fill(r: int, recs: list[dict], new: bool = False) -> dict:
        main = min(recs, key=lambda x: (x.get("include") != "yes", ORDER[x["media_type"]])) if recs else None
        m = analyser(main) if main else {}
        checks = [ADDED] if new else []
        checks += [x.get("check") for x in recs if x.get("check")]
        if m.get("note"):
            checks.append(m["note"])
        values = {
            "In data/input": presence(recs),
            "Files": "\n".join(f for x in recs for f in (x.get("video_file"), x.get("transcript_file")) if f),
            "Transcript": "\n".join(dict.fromkeys(transcript_info(x) for x in recs)),
            "Slides": "\n".join(slides_info(x) for x in recs if x["media_type"] == "video") or ("-" if recs else ""),
            "Minutes": minutes(recs),
            "Topics (from the recording)": "\n".join(t["title"] for t in m.get("topics", [])),
            "Check": "\n".join(dict.fromkeys(checks)),
        }
        for name, value in values.items():
            ws.cell(r, col[name], value).alignment = wrap
        state = "new" if new else "missing" if not recs else "ok"
        ws.cell(r, col["In data/input"]).fill = PatternFill("solid", fgColor=FILL[state])
        return m

    counts = {"rows": 0, "found": 0, "missing": [], "added": [], "no_transcript": 0, "analysed": 0}
    for r in range(head + 1, ws.max_row + 1):
        eid = inventory._norm_id(ws.cell(r, cols["id"]).value)
        if not eid:
            continue
        recs = by_id.get(eid, [])
        m = fill(r, recs)
        counts["rows"] += 1
        counts["analysed"] += bool(m.get("topics"))
        if recs:
            counts["found"] += 1
        else:
            counts["missing"].append(f"{eid} {ws.cell(r, cols['name']).value or ''}".strip())

    used = set(entries)
    for rec in sorted(added, key=lambda x: (ORDER[x["media_type"]], x.get("title", "").lower())):
        r = ws.max_row + 1
        number = inventory.LEADING_NUMBER.match(Path(rec.get("video_file") or rec.get("transcript_file")).stem)
        eid = inventory._norm_id(number.group(1)) if number else ""
        ws.cell(r, cols["id"], "" if eid in used else (int(eid) if eid.isdigit() else eid))
        ws.cell(r, cols["name"], rec.get("title", ""))
        m = fill(r, [rec], new=True)
        ws.cell(r, cols["description"], m.get("summary", "")).alignment = wrap
        for c in range(1, col["Check"] + 1):
            if c != col["In data/input"]:
                ws.cell(r, c).fill = PatternFill("solid", fgColor=FILL["new"])
        counts["added"].append(rec.get("title", ""))
        counts["analysed"] += bool(m.get("topics"))
    counts["no_transcript"] = sum(store.transcript_path(x) is None for x in rows)

    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = config.OUTPUT_DIR / f"{path.stem}_checked.xlsx"
    try:
        wb.save(out)
    except PermissionError:   # the previous copy is open in Excel
        out = out.with_name(f"{path.stem}_checked_{datetime.now():%H%M%S}.xlsx")
        wb.save(out)
    return out, counts


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-analyse", action="store_true", help="no model calls: only what is there (free)")
    args = ap.parse_args()
    config.setup_logging()
    try:
        out, c = build(analyse=not args.no_analyse)
    except (FileNotFoundError, RuntimeError, ValueError) as e:
        log.error("%s", e)
        return 1
    print(f"\n{c['found']} of {c['rows']} rows of the Excel have a recording or transcript in data/input.")
    if c["missing"]:
        print(f"MISSING ({len(c['missing'])}): " + "; ".join(c["missing"]))
    if c["added"]:
        print(f"Added as new rows ({len(c['added'])}): " + "; ".join(c["added"]))
    if c["no_transcript"]:
        print(f"{c['no_transcript']} recordings have no transcript yet: uv run pipeline.py --mode all")
    print(f"Analysed (topics, description): {c['analysed']}")
    print(f"\nWritten: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
