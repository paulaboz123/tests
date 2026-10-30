"""description.xlsx checked against data/input (catalog.py): what is there, what is missing,
recordings the Excel doesn't list added with a description from their transcript and slides."""
from __future__ import annotations

from openpyxl import Workbook, load_workbook
from test_end_to_end import _prepare_one_recording

import catalog
import config
import store


def test_the_description_excel_is_checked_and_completed(gateway):
    _prepare_one_recording()                                   # "13. Joining data.vtt", processed
    wb = Workbook()
    ws = wb.active
    ws.append(["Video catalogue"])
    ws.append(["ID", "Name", "Description", "Owner"])
    ws.append([13, "Joining data", "How to join two tables.", "Ann"])
    ws.append([14, "Pivot tables", "Pivots.", ""])
    wb.save(config.INPUT_DIR / "description.xlsx")
    rows = store.read_manifest() + [
        {"video_id": "weekly-sync-1234", "include": "yes", "title": "Weekly sync", "media_type": "video",
         "video_file": "Weekly sync.mp4", "transcript_file": "", "check": "NO TRANSCRIPT"},
        {"video_id": "weekly-sync-audio", "include": "no", "title": "Weekly sync", "media_type": "audio",
         "video_file": "Weekly sync.m4a", "transcript_file": "", "check": "AUDIO OF A VIDEO ALREADY LISTED"}]
    store.write_manifest(rows)
    out, c = catalog.build()
    assert c["found"] == 1 and c["missing"] == ["14 Pivot tables"] and c["added"] == ["Weekly sync"]
    assert c["analysed"] == 1 and c["no_transcript"] == 2
    ws = load_workbook(out).active
    head = [x.value for x in ws[2]]
    assert head[:4] == ["ID", "Name", "Description", "Owner"] and head[4:] == catalog.COLUMNS   # theirs kept
    col = {n: head.index(n) for n in head}
    r13, r14, new = ([x.value for x in ws[i]] for i in (3, 4, 5))
    assert r13[col["In data/input"]] == "transcript only" and r13[col["Transcript"]] == ".vtt (given)"
    assert r13[col["Topics (from the recording)"]] == "Join keys\nChecking a join" and r13[3] == "Ann"
    assert r14[col["In data/input"]] == "MISSING" and r14[col["Description"]] == "Pivots."
    assert new[col["Name"]] == "Weekly sync" and new[col["In data/input"]] == "video" and new[0] is None
    assert "ADDED" in new[col["Check"]] and "not processed yet" in new[col["Check"]]
    assert new[col["Transcript"]].startswith("NONE yet") and new[col["Slides"]].startswith("not read yet")
    assert ws.max_row == 5                                     # the audio copy of the video is not a new row
    calls = len(gateway.state["calls"])
    catalog.build()
    assert len(gateway.state["calls"]) == calls                # the analysis is cached


def test_new_recording_gets_its_number_and_description(gateway):
    _prepare_one_recording()
    wb = Workbook()
    wb.active.append(["id", "name"])
    wb.active.append([5, "Something else"])
    wb.save(config.INPUT_DIR / "description.xlsx")
    out, c = catalog.build()
    ws = load_workbook(out).active
    head = [x.value for x in ws[1]]
    new = [x.value for x in ws[3]]
    assert c["added"] == ["Joining data"] and new[0] == 13 and head[2] == "description"
    assert new[2] == "How to join tables." and new[head.index("In data/input")] == "transcript only"
    _, free = catalog.build(analyse=False)
    assert free["analysed"] == 0
