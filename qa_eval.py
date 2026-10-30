"""
qa_eval.py - measures how good the text answers are (answer.py), on a set of test questions:
does it find the right information, is every statement true to the material (no hallucination),
is it complete and on the point, does it decline what the material doesn't cover, is it written
in own words, and does every sentence say where it comes from.

  uv run qa_eval.py --draft 20        # writes eval/qa_set.csv from your material: check and edit it in Excel
  uv run qa_eval.py                   # answers every question and scores it -> data/output/<time>-qa-eval/
  uv run qa_eval.py --compare         # the same set with each retrieval mode, side by side (before / after)
  uv run qa_eval.py --limit 5         # the first 5 questions; --ids q01,q07 for chosen ones
  uv run qa_eval.py --model <id> --judge <id>    # writer and judge models (default CHAT_MODEL, JUDGE_MODEL)
  uv run qa_eval.py --set "Video_Accuracy_Test_Plan.xlsx"    # the client's test plan, as it is (see below)

eval/qa_set.csv, one row per question (add your own and the client's questions any time):
  id                q01, q02 ...
  type              general | specific | out_of_scope (the material does not answer it: the right
                    result is "not covered", a test for hallucination)
  question          as a learner would ask it
  expected_answer   what a good answer says (optional; then correctness is scored)
  key_facts         facts a good answer must contain, separated by ';' (optional; the strongest check
                    that nothing important is lost)
  expected_sources  where the answer is: a title (or part of it), optionally @time, separated by ';'
                    e.g. "Scrum basics@03:10; Scrum Guide" (optional; then retrieval is scored)
  length            short | medium | long (optional, default short)
  notes

A test plan in Excel (.xlsx) is read as it is: the header row is found by its names (Test ID,
Category, Prompt, Source Video (ID+name), Expected fact; other columns are ignored). Rows with the
same Test ID (or the same prompt) are one question with several expected videos and facts; empty
merged cells take the value above. "4 - Tools @SE" is matched by its number 4 (catalog_id in
description.xlsx) or by its name. Categories such as "out of scope", "negative" or "hallucination"
expect the answer "not covered". The results are also written into a copy of the plan
(<name>_results.xlsx): the answer, its sources (video, time), whether each row's expected fact is in
the answer and its video was found, claims not in the sources, and the verdict.

Result folder: qa_report.html and qa_report.md (scorecard, failures with reasons, every answer with
its sources and timestamps, how each metric is measured and why), qa_results.csv (one row per
question, with empty columns for your own verdict), and each answer's own folder.
"""
from __future__ import annotations

import argparse
import csv
import html
import logging
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import answer as answer_mod
import config
import llm
import retrieval
import store

log = logging.getLogger("qa_eval")

def qa_set_path() -> Path:
    return config.EVAL_DIR / "qa_set.csv"
SET_FIELDS = ["id", "type", "question", "expected_answer", "key_facts", "expected_sources", "length", "notes"]
TYPES = ("general", "specific", "out_of_scope")


# --- The judges ---------------------------------------------------------------------------------------
GROUNDING = """You check every claim of an answer against the excerpts it cites, strictly. Split each sentence into
its separate factual claims. For each claim decide, using ONLY the cited excerpts (not your own knowledge):
supported = the excerpts state it (a paraphrase is fine); partial = only part of it is stated, or it is stated more
strongly than the excerpts do; unsupported = the excerpts do not say it (even if it is true in general);
contradicted = the excerpts say something different. Give a short reason for anything not supported."""

GROUNDING_SCHEMA = {"type": "object", "properties": {"claims": {"type": "array", "items": {"type": "object", "properties": {
    "sentence": {"type": "integer"}, "claim": {"type": "string"},
    "verdict": {"type": "string", "enum": ["supported", "partial", "unsupported", "contradicted"]},
    "reason": {"type": "string"}}, "required": ["sentence", "claim", "verdict"]}}}, "required": ["claims"]}

GRADING = """You grade an answer to a learner's question, strictly and fairly. You get the question, the answer, the
excerpts the system found in the training material, and (when there is one) the expected answer and key facts.
Score 1-5 (5 best):
relevance     it answers exactly what was asked, with the most important information first, nothing off the point
completeness  it contains the most important information that the excerpts give for this question
correctness   it agrees with the expected answer (0 if there is no expected answer)
clarity       clear, well structured, easy for a learner
For each key fact say: yes (contained, in any wording), partial, or no."""

GRADING_SCHEMA = {"type": "object", "properties": {
    "relevance": {"type": "integer"}, "completeness": {"type": "integer"}, "correctness": {"type": "integer"},
    "clarity": {"type": "integer"},
    "facts": {"type": "array", "items": {"type": "object", "properties": {
        "fact": {"type": "string"}, "found": {"type": "string", "enum": ["yes", "partial", "no"]}},
        "required": ["fact", "found"]}},
    "comment": {"type": "string"}}, "required": ["relevance", "completeness", "correctness", "clarity", "facts"]}

POINTS = {"supported": 1.0, "partial": 0.5, "unsupported": 0.0, "contradicted": 0.0}


def judge_model(args_judge: str | None = None) -> str:
    return args_judge or config.JUDGE_MODEL or config.REVIEW_MODEL or config.CHAT_MODEL


def grounding(rec: dict, model: str) -> dict:
    """Claim by claim: is it in the excerpts the sentence cites?"""
    parts = []
    for n, s in enumerate(rec["sentences"], 1):
        cited = "\n".join(f"  [{x}] {rec['excerpts'][x]['text']}" for x in s["sources"]) or "  (no source cited)"
        parts.append(f"SENTENCE {n}: {s['text']}\nCITED EXCERPTS:\n{cited}")
    out = llm.chat_json(f"Question: {rec['question']}\n\n" + "\n\n".join(parts), GROUNDING_SCHEMA, system=GROUNDING,
                        step="judge grounding", model=model, temperature=0.0, max_tokens=3000)
    claims = [c for c in out.get("claims", []) if c.get("verdict") in POINTS]
    n = len(claims)
    bad = [c for c in claims if c["verdict"] in ("unsupported", "contradicted")]
    return {"claims": claims, "n_claims": n,
            "faithfulness": round(sum(POINTS[c["verdict"]] for c in claims) / n, 3) if n else None,
            "hallucination_rate": round(len(bad) / n, 3) if n else None,
            "unsupported": [f"{c['claim']} ({c['verdict']}: {c.get('reason', '')})" for c in bad]}


def grading(rec: dict, row: dict, model: str) -> dict:
    facts = row.get("_facts") or [f.strip() for f in (row.get("key_facts") or "").split(";") if f.strip()]
    excerpts = "\n".join(f"[{k}] {e['text']}" for k, e in rec["excerpts"].items())
    prompt = (f"Question: {rec['question']}\n\nAnswer:\n{rec['text']}\n\nExcerpts found:\n{excerpts}\n\n"
              f"Expected answer: {row.get('expected_answer') or '(none)'}\n"
              f"Key facts:\n" + ("\n".join(f"- {f}" for f in facts) or "(none)"))
    out = llm.chat_json(prompt, GRADING_SCHEMA, system=GRADING, step="judge grading", model=model, temperature=0.0,
                        max_tokens=1500)
    found = {"yes": 1.0, "partial": 0.5, "no": 0.0}
    judged = out.get("facts", [])
    per_fact = []
    for i, fact in enumerate(facts):   # the judge's verdict for each fact: same text, else same position
        hit = next((f for f in judged if (f.get("fact") or "").strip().lower() == fact.lower()), None)
        hit = hit or (judged[i] if i < len(judged) else {})
        per_fact.append((fact, hit.get("found", "no") if hit.get("found") in found else "no"))
    got = [found[v] for _, v in per_fact]
    clamp = lambda v: max(0, min(5, int(v or 0)))  # noqa: E731
    return {"relevance": clamp(out.get("relevance")), "completeness": clamp(out.get("completeness")),
            "correctness": clamp(out.get("correctness")) if row.get("expected_answer") else None,
            "clarity": clamp(out.get("clarity")),
            "key_fact_recall": round(sum(got) / len(facts), 3) if facts else None,
            "missing_facts": [f for f, v in per_fact if v == "no"], "fact_results": per_fact,
            "comment": out.get("comment", "")}


# --- Retrieval scores (no model) ------------------------------------------------------------------------
def parse_sources(spec: str) -> list[tuple[str, float | None]]:
    """'Scrum basics@03:10; Scrum Guide' -> [('scrum basics', 190.0), ('scrum guide', None)]."""
    out = []
    for part in (spec or "").split(";"):
        title, _, at = part.strip().partition("@")
        if not title.strip():
            continue
        seconds = None
        if at.strip():
            nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", at)]
            seconds = sum(v * 60 ** i for i, v in enumerate(reversed(nums))) if nums else None
        out.append((title.strip().lower(), seconds))
    return out


def matches(seg: dict, title: str, seconds: float | None) -> bool:
    if title not in (seg.get("title", "") or "").lower() and title != seg.get("video_id", "").lower():
        return False
    return seconds is None or (seg.get("start", 0) - 45 <= seconds <= seg.get("end", 0) + 45)


def same_video(seg: dict, cid: str, name: str) -> bool:
    """Clara's "4 - Tools @SE": the recording with catalog_id 4, or with that name in its title."""
    if cid and str(seg.get("catalog_id", "")) == cid:
        return True
    a, b = set(re.findall(r"[a-z0-9]+", name.lower())), set(re.findall(r"[a-z0-9]+", (seg.get("title") or "").lower()))
    return bool(a) and len(a & b) / len(a) >= 0.75


def parse_video(spec: str) -> tuple[str, str]:
    """'4 - Tools @SE' -> ('4', 'Tools @SE'); 'Tools @SE' -> ('', 'Tools @SE')."""
    m = re.match(r"\s*#?(\d+)\s*[-.:)]?\s*(.*)$", spec or "")
    return (m.group(1), m.group(2).strip()) if m else ("", (spec or "").strip())


def retrieval_scores(rec: dict, row: dict) -> dict:
    """Did the search find where the answer is? hit = among the excerpts the answer was written
    from (every expected video, when there are several); rank = place of the first right excerpt
    among all candidates (1 = first); cited = the answer cites it. Per expected video too."""
    segs = store.all_segments()
    if row.get("_sources"):   # a test plan: whole videos, by number or name
        tests = [(f"{cid} - {name}".strip(" -"), lambda seg, c=cid, n=name: same_video(seg, c, n)) for cid, name in row["_sources"]]
    else:
        tests = [(f"{t}" + (f"@{store.fmt_ts(s)}" if s is not None else ""), lambda seg, t=t, s=s: matches(seg, t, s))
                 for t, s in parse_sources(row.get("expected_sources", ""))]
    if not tests:
        return {"retrieval_hit": None, "reciprocal_rank": None, "cited_expected": None, "source_results": []}
    kept = [i for i in rec.get("kept", []) if i in segs]
    cited_ids = {rec["excerpts"][x]["id"] for s in rec.get("sentences", []) for x in s["sources"] if x in rec.get("excerpts", {})}
    per = [{"source": name, "found": any(test(segs[i]) for i in kept),
            "cited": any(test(segs[i]) for i in cited_ids if i in segs)} for name, test in tests]
    ok = lambda sid: sid in segs and any(test(segs[sid]) for _, test in tests)  # noqa: E731
    rank = next((i for i, c in enumerate(rec.get("candidates", []), 1) if ok(c["id"])), None)
    return {"retrieval_hit": round(sum(p["found"] for p in per) / len(per), 3),
            "reciprocal_rank": round(1 / rank, 3) if rank else 0.0,
            "cited_expected": round(sum(p["cited"] for p in per) / len(per), 3) if rec.get("sentences") else 0.0,
            "source_results": per}


# --- One question -------------------------------------------------------------------------------------
def evaluate(row: dict, mode: str, model: str | None, judge: str, use_judge: bool = True) -> dict:
    t0 = time.perf_counter()
    rec = answer_mod.answer(row["question"], row.get("length") or "short", "beginner", mode, model)
    answered = not rec["not_in_material"]
    out_of_scope = row.get("type") == "out_of_scope"
    c = rec.get("checks") or {}
    res = {"id": row["id"], "type": row.get("type", ""), "category": row.get("category", ""), "question": row["question"],
           "mode": rec.get("mode", mode),
           "answered": answered, "abstain_correct": answered != out_of_scope, "answer": rec["text"],
           "sources": "; ".join(f"[{s['n']}] {s['title']} · {s['where']}" for s in rec.get("sources", [])),
           "evidence_verified": c.get("evidence_verified"), "unsourced_sentences": c.get("unsourced_sentences", 0),
           "invalid_labels": len(c.get("invalid_labels", [])), "numbers_not_in_sources": len(c.get("numbers_not_in_sources", [])),
           "copy_rate": c.get("copy_rate"), "longest_copy_words": c.get("longest_copy_words"), "words": c.get("words", 0),
           "folder": Path(rec.get("folder", "")).name, **retrieval_scores(rec, row)}
    mark = llm.meter_mark()
    if answered and use_judge:
        try:
            g = grounding(rec, judge)
            res.update({k: g[k] for k in ("n_claims", "faithfulness", "hallucination_rate")},
                       unsupported_claims=" | ".join(g["unsupported"]))
            res.update(grading(rec, row, judge))
        except llm.LLMError as e:
            res["judge_error"] = str(e)[:200]
    judge_cost = sum(x["cost_usd"] for x in llm.calls_since(mark))
    res.update(cost_usd=round(rec["cost_usd"], 5), judge_cost_usd=round(judge_cost, 5),
               seconds=round(time.perf_counter() - t0, 1))
    res["passed"], res["why"] = verdict(res, row)
    return res


def verdict(r: dict, row: dict) -> tuple[bool, str]:
    """Pass rules (stated in the report): out of scope -> declined; in scope -> answered, no
    unsupported or contradicted claim, evidence found, no invented source or number, key facts
    and relevance good enough."""
    if row.get("type") == "out_of_scope":
        return (not r["answered"], "" if not r["answered"] else "answered a question the material does not cover")
    if not r["answered"]:
        return False, "declined, but the material should cover it"
    why = []
    if r.get("hallucination_rate"):
        why.append(f"{r['hallucination_rate']:.0%} of claims not in the sources")
    if r.get("evidence_verified") is not None and r["evidence_verified"] < 0.8:
        why.append(f"evidence found for only {r['evidence_verified']:.0%} of sentences")
    if r["invalid_labels"] or r["unsourced_sentences"]:
        why.append("sentences without a real source")
    if r["numbers_not_in_sources"]:
        why.append("numbers not in the sources")
    if r.get("key_fact_recall") is not None and r["key_fact_recall"] < 0.6:
        why.append(f"key facts {r['key_fact_recall']:.0%}")
    if r.get("relevance") is not None and r["relevance"] < 4:
        why.append(f"relevance {r['relevance']}/5")
    if r.get("retrieval_hit") is not None and r["retrieval_hit"] < 1:
        missed = [p["source"] for p in r.get("source_results", []) if not p["found"]]
        why.append("expected video not found: " + ", ".join(missed) if missed else "the expected source was not found")
    return not why, "; ".join(why)


# --- Summary -------------------------------------------------------------------------------------------
def mean(rows: list[dict], key: str) -> float | None:
    vals = [float(r[key]) for r in rows if r.get(key) not in (None, "")]
    return round(sum(vals) / len(vals), 3) if vals else None


def summary(rows: list[dict]) -> dict:
    inscope = [r for r in rows if r["type"] != "out_of_scope"]
    oos = [r for r in rows if r["type"] == "out_of_scope"]
    answered = [r for r in inscope if r["answered"]]
    return {"questions": len(rows), "passed": sum(r["passed"] for r in rows),
            "pass_rate": round(sum(r["passed"] for r in rows) / len(rows), 3) if rows else None,
            "retrieval_hit": mean(inscope, "retrieval_hit"), "mrr": mean(inscope, "reciprocal_rank"),
            "cited_expected": mean(inscope, "cited_expected"),
            "faithfulness": mean(answered, "faithfulness"), "hallucination_rate": mean(answered, "hallucination_rate"),
            "answers_with_hallucination": sum(1 for r in answered if r.get("hallucination_rate")),
            "evidence_verified": mean(answered, "evidence_verified"),
            "key_fact_recall": mean(answered, "key_fact_recall"), "correctness": mean(answered, "correctness"),
            "relevance": mean(answered, "relevance"), "completeness": mean(answered, "completeness"),
            "clarity": mean(answered, "clarity"),
            "declined_out_of_scope": mean([{"x": float(not r["answered"])} for r in oos], "x"),
            "false_refusals": mean([{"x": float(not r["answered"])} for r in inscope], "x"),
            "copy_rate": mean(answered, "copy_rate"),
            "invalid_or_unsourced": sum(r["invalid_labels"] + r["unsourced_sentences"] for r in answered),
            "numbers_not_in_sources": sum(r["numbers_not_in_sources"] for r in answered),
            "cost_per_answer_usd": mean(rows, "cost_usd"), "judge_cost_usd": round(sum(r["judge_cost_usd"] for r in rows), 4),
            "seconds_per_answer": mean(rows, "seconds")}


METRICS = [   # key, area, name, how it is measured, why it matters, target, higher is better
    ("retrieval_hit", "Finding the information", "Retrieval hit rate",
     "share of questions where an excerpt from the expected source (and time) is among those the answer is written from",
     "if the right excerpt is not found, no model can answer well: this separates search problems from writing problems",
     ">= 90%", True),
    ("mrr", "Finding the information", "Mean reciprocal rank (MRR)",
     "1 / position of the first right excerpt among all candidates, averaged (1.0 = always first)",
     "how high the right information ranks; low values mean it is found by luck", ">= 0.7", True),
    ("cited_expected", "Finding the information", "Answer cites the expected source",
     "share of answers that cite an excerpt from the expected source", "the right source was not only found but used", ">= 85%", True),
    ("faithfulness", "No hallucination", "Faithfulness (claims supported)",
     "the judge model splits each answer into single claims and checks each one against the excerpts that sentence cites "
     "(supported 1, partial 0.5, unsupported or contradicted 0); average per answer",
     "the core measure that the answer says only what the material says", ">= 95%", True),
    ("hallucination_rate", "No hallucination", "Hallucination rate",
     "share of claims that the cited excerpts do not contain or contradict", "every such claim could teach something wrong",
     "0%", False),
    ("answers_with_hallucination", "No hallucination", "Answers with any unsupported claim",
     "number of answers with at least one unsupported or contradicted claim", "a single wrong claim spoils an answer", "0", False),
    ("evidence_verified", "No hallucination", "Evidence found word for word",
     "every sentence gives a short quote from its source; share of quotes found in that source (checked in code, no model)",
     "a check of the sources that does not depend on a model's judgement, and lets a person verify in seconds", ">= 95%", True),
    ("numbers_not_in_sources", "No hallucination", "Numbers not in the sources",
     "numbers in answers that the cited excerpts do not contain (code check; number words count too)",
     "invented numbers are the most damaging and the easiest to check", "0", False),
    ("invalid_or_unsourced", "No hallucination", "Invented or missing sources",
     "source labels the model invented, plus sentences that cite nothing (code check)", "every statement must be traceable",
     "0", False),
    ("key_fact_recall", "Correct and complete", "Key-fact recall",
     "share of the key facts of the test question that the answer contains, in any wording (judge model)",
     "nothing important is lost; the strongest check of completeness", ">= 80%", True),
    ("correctness", "Correct and complete", "Correctness (1-5)",
     "agreement with the expected answer (judge model; only questions with an expected answer)", "", ">= 4", True),
    ("relevance", "Correct and complete", "Relevance (1-5)",
     "answers exactly what was asked, most important first (judge model)", "people get what they need, not a summary of everything", ">= 4", True),
    ("completeness", "Correct and complete", "Completeness (1-5)",
     "contains the most important information the found excerpts give (judge model)", "", ">= 4", True),
    ("clarity", "Correct and complete", "Clarity (1-5)", "clear and well structured for a learner (judge model)", "", ">= 4", True),
    ("declined_out_of_scope", "Knows what it doesn't know", "Declines out-of-scope questions",
     "share of out-of-scope questions answered with 'not covered' instead of an answer",
     "the clearest hallucination test: it must not answer from general knowledge", "100%", True),
    ("false_refusals", "Knows what it doesn't know", "False refusals",
     "share of answerable questions it declined", "too careful is also a failure", "<= 5%", False),
    ("copy_rate", "Own words", "Copied word for word",
     "share of the answer's 8-word sequences found word for word in the excerpts (code check)",
     "the answer should give the most important information in new sentences, not repeat the input", "<= 20%", False),
    ("cost_per_answer_usd", "Cost and time", "Cost per answer (USD)", "model calls of the answer (search, rerank, writing)", "", "", False),
    ("seconds_per_answer", "Cost and time", "Seconds per answer", "", "", "", False),
]


def fmt(key: str, v) -> str:
    if v is None:
        return "-"
    if key in ("retrieval_hit", "cited_expected", "faithfulness", "hallucination_rate", "evidence_verified",
               "key_fact_recall", "declined_out_of_scope", "false_refusals", "copy_rate", "pass_rate"):
        return f"{v:.0%}"
    if key == "cost_per_answer_usd":
        return f"${v:.4f}"
    return f"{v:g}" if isinstance(v, float) else str(v)


# --- Report ---------------------------------------------------------------------------------------------
def report(folder: Path, runs: dict[str, list[dict]], settings: dict, previous: dict | None) -> None:
    sums = {m: summary(rows) for m, rows in runs.items()}
    modes = list(runs)
    md = [f"# Answer quality: {settings['questions']} questions", "",
          f"{settings['created']} · writer {settings['writer']} · judge {settings['judge']} · rerank {settings['rerank']} · "
          f"set {settings['set']}", ""]
    if settings["writer"] == settings["judge"]:
        md += ["*Note: the writer and the judge are the same model. For an independent grade set JUDGE_MODEL to "
               "another model (e.g. Claude Opus when answers are written by Sonnet).*", ""]
    md += ["## Scorecard", "", "| Area | Metric | " + " | ".join(modes) + " | Target |",
           "|---|---|" + "---|" * len(modes) + "---|"]
    md.append("| Overall | **Questions passed** | " + " | ".join(
        f"**{sums[m]['passed']}/{sums[m]['questions']} ({fmt('pass_rate', sums[m]['pass_rate'])})**" for m in modes) + " | |")
    for key, area, name, _how, _why, target, _up in METRICS:
        md.append(f"| {area} | {name} | " + " | ".join(fmt(key, sums[m][key]) for m in modes) + f" | {target} |")
    if previous:
        md += ["", "Previous run (same set, " + previous.get("created", "") + "): passed "
               f"{previous.get('passed')}/{previous.get('questions')}, faithfulness {fmt('faithfulness', previous.get('faithfulness'))}, "
               f"retrieval hit {fmt('retrieval_hit', previous.get('retrieval_hit'))}."]
    best = modes[-1]
    rows = runs[best]
    md += ["", f"## By question type ({best})", "", "| Type | Questions | Passed | Faithfulness | Key facts | Relevance |",
           "|---|---|---|---|---|---|"]
    for t in TYPES:
        sel = [r for r in rows if r["type"] == t]
        if sel:
            s = summary(sel)
            md.append(f"| {t} | {len(sel)} | {s['passed']} | {fmt('faithfulness', s['faithfulness'])} | "
                      f"{fmt('key_fact_recall', s['key_fact_recall'])} | {fmt('relevance', s['relevance'])} |")
    cats = list(dict.fromkeys(r.get("category") for r in rows if r.get("category")))
    if cats:
        md += ["", f"## By category ({best})", "", "| Category | Questions | Passed | Retrieval hit | Faithfulness | Key facts |",
               "|---|---|---|---|---|---|"]
        for cat in cats:
            s = summary([r for r in rows if r.get("category") == cat])
            md.append(f"| {cat} | {s['questions']} | {s['passed']} | {fmt('retrieval_hit', s['retrieval_hit'])} | "
                      f"{fmt('faithfulness', s['faithfulness'])} | {fmt('key_fact_recall', s['key_fact_recall'])} |")
    failed = [r for r in rows if not r["passed"]]
    md += ["", f"## Questions to look at ({len(failed)})", ""]
    for r in failed:
        md += [f"**{r['id']} ({r['type']}): {r['question']}**", f"- Why: {r['why']}"]
        if r.get("unsupported_claims"):
            md.append(f"- Claims not in the sources: {r['unsupported_claims']}")
        if r.get("missing_facts"):
            md.append(f"- Missing key facts: {'; '.join(r['missing_facts'])}")
        md += [f"- Answer: {r['answer']}", f"- Sources: {r['sources'] or '-'}", ""]
    md += ["## All answers", ""]
    for r in rows:
        md += [f"**{r['id']}: {r['question']}** ({'passed' if r['passed'] else 'look at it'})", "", r["answer"], "",
               f"Sources: {r['sources'] or '-'}", ""]
    md += ["## How each metric is measured, and why", "", "| Metric | How | Why | Target |", "|---|---|---|---|"]
    md += [f"| {name} | {how} | {why} | {target} |" for key, _a, name, how, why, target, _u in METRICS if how]
    md += ["", "**Pass rule per question.** Out of scope: declined. In scope: answered; no unsupported or contradicted "
           "claim; evidence found for at least 80% of sentences; no sentence without a real source; no number missing "
           "from the sources; key-fact recall at least 60% and relevance at least 4/5 (when scored); and the expected "
           "source found (when given).", "",
           "**Why a model as judge, and its limits.** Claim-by-claim checking is what a careful reviewer does; a model "
           "does it for every answer at the cost of cents, with a reason for every claim it rejects. It can still err, "
           "so the code checks (evidence quotes, numbers, sources, copying) do not depend on a model, the judge is best a "
           "different model than the writer, and qa_results.csv has columns for your own verdict on a sample.", ""]
    (folder / "qa_report.md").write_text("\n".join(md), encoding="utf-8")
    (folder / "qa_report.html").write_text(html_page(md), encoding="utf-8")


def html_page(md: list[str]) -> str:
    """The markdown report as a simple page (tables, headings, bold), without extra packages."""
    out, table = [], []

    def inline(t: str) -> str:
        t = html.escape(t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
        return re.sub(r"\*(.+?)\*", r"<i>\1</i>", t)

    def flush():
        if table:
            rows = [r for r in table if not re.fullmatch(r"\|[-| ]+\|", r)]
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
            out.append("<table><thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in cells[0]) + "</tr></thead><tbody>"
                       + "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in cells[1:])
                       + "</tbody></table>")
            table.clear()
    for line in md:
        if line.startswith("|"):
            table.append(line)
            continue
        flush()
        if line.startswith("# "):
            out.append(f"<h1>{inline(line[2:])}</h1>")
        elif line.startswith("## "):
            out.append(f"<h2>{inline(line[3:])}</h2>")
        elif line.startswith("- "):
            out.append(f"<p class='li'>• {inline(line[2:])}</p>")
        elif line.strip():
            out.append(f"<p>{inline(line)}</p>")
    flush()
    css = ("body{font:15px/1.5 'Segoe UI',Arial,sans-serif;max-width:1200px;margin:24px auto;padding:0 16px;color:#14232a}"
           "table{border-collapse:collapse;width:100%;margin:8px 0 16px}th,td{border:1px solid #d9e2e3;padding:6px 8px;"
           "text-align:left;vertical-align:top}th{background:#eef3f4}h1{color:#0b6e79}h2{color:#0b6e79;margin-top:28px}"
           ".li{margin:2px 0 2px 12px}")
    return f"<!doctype html><meta charset='utf-8'><title>Answer quality</title><style>{css}</style>{''.join(out)}"


RESULT_FIELDS = ["id", "type", "category", "mode", "question", "passed", "why", "answered", "abstain_correct", "retrieval_hit",
                 "reciprocal_rank", "cited_expected", "n_claims", "faithfulness", "hallucination_rate", "unsupported_claims",
                 "evidence_verified", "numbers_not_in_sources", "invalid_labels", "unsourced_sentences", "key_fact_recall",
                 "missing_facts", "correctness", "relevance", "completeness", "clarity", "copy_rate", "longest_copy_words",
                 "words", "cost_usd", "judge_cost_usd", "seconds", "answer", "sources", "folder", "judge_error",
                 "your_verdict", "comments"]


# --- Test set ---------------------------------------------------------------------------------------------
OUT_OF_SCOPE = re.compile(r"out[ -]?of[ -]?scope|negative|not (in|covered)|unanswerable|no answer|halluc|trap|outside|"
                          r"should (decline|refuse|say)", re.IGNORECASE)
HEADERS = {"id": ("test id", "test case", "id", "tc"), "category": ("category", "type"), "question": ("prompt", "question"),
           "source": ("source video", "source", "video"), "fact": ("expected fact", "expected facts", "expected result",
                                                                    "expected"), "answer": ("expected answer",)}


def _header_map(cells: list) -> dict[str, int]:
    """Column of each field in a header row, by its name ('Source Video (ID+name)' -> source)."""
    names = [str(c or "").strip().lower() for c in cells]
    out: dict[str, int] = {}
    for field, keys in HEADERS.items():
        for key in keys:   # most specific name first
            col = next((i for i, n in enumerate(names) if n and (n == key or n.startswith(key)) and i not in out.values()), None)
            if col is not None:
                out[field] = col
                break
    return out


def load_test_plan(path: Path) -> list[dict]:
    """The client's test plan (.xlsx), as it is: one question per Test ID (or prompt), with all its
    expected videos and facts, and where each row is (to write the results back into a copy)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    tests: dict[str, dict] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        head = next(((i, _header_map(list(r))) for i, r in enumerate(rows[:20]) if "question" in _header_map(list(r))), None)
        if not head:
            continue
        h, cols = head
        def get(r, f, cols=cols):
            return str(r[cols[f]]).strip() if f in cols and cols[f] < len(r) and r[cols[f]] is not None else ""
        last = {"id": "", "category": "", "question": ""}
        for i, r in enumerate(rows[h + 1:], h + 2):
            vals = {f: get(r, f) for f in ("id", "category", "question", "source", "fact", "answer")}
            if not any(vals.values()):
                continue
            for f in ("id", "category", "question"):   # merged cells: the value is in the first row only
                if vals[f]:
                    last[f] = vals[f]
                elif vals["source"] or vals["fact"]:
                    vals[f] = last[f]
            if not vals["question"]:
                continue
            key = vals["id"] or vals["question"]
            t = tests.setdefault(key, {"id": vals["id"] or f"q{len(tests) + 1:02d}", "category": vals["category"],
                                       "question": vals["question"], "expected_answer": vals["answer"], "_facts": [],
                                       "_sources": [], "_rows": [], "length": "medium", "notes": f"{path.name} / {ws.title}"})
            if vals["fact"] and vals["fact"] not in t["_facts"]:
                t["_facts"].append(vals["fact"])
            if vals["source"] and vals["source"].lower() not in ("-", "n/a", "na", "none", "no video"):
                video = parse_video(vals["source"])
                if video not in t["_sources"]:
                    t["_sources"].append(video)
            t["_rows"].append({"sheet": ws.title, "row": i, "header": h + 1, "source": vals["source"], "fact": vals["fact"]})
    wb.close()
    out = list(tests.values())
    for t in out:
        oos = OUT_OF_SCOPE.search(t["category"]) or (not t["_sources"] and any(OUT_OF_SCOPE.search(f) for f in t["_facts"]))
        t["type"] = "out_of_scope" if oos else ("general" if len(t["_sources"]) != 1 else "specific")
        t["key_facts"] = " ; ".join(t["_facts"])
        t["expected_sources"] = " ; ".join(f"{c} - {n}".strip(" -") for c, n in t["_sources"])
    if not out:
        raise ValueError(f"No test plan found in {path.name}: no sheet has a header row with a 'Prompt' (or 'Question') column")
    return out


def write_back(plan: Path, rows: list[dict], results: list[dict], out: Path) -> Path:
    """A copy of the test plan with the results next to each row (new columns on the right)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(plan)
    by_id = {r["id"]: r for r in results}
    cols = ["System answer", "Sources used (video · time)", "Expected fact in the answer", "Expected video found",
            "Claims not in the sources", "Faithfulness", "Result", "Why"]
    done_sheets: dict[str, int] = {}
    for t in rows:
        res = by_id.get(t["id"])
        if not res:
            continue
        facts = dict(res.get("fact_results") or [])
        videos = {p["source"]: p for p in res.get("source_results") or []}
        for place in t.get("_rows", []):
            ws = wb[place["sheet"]]
            if place["sheet"] not in done_sheets:   # the header row of the new columns
                head_row = place["header"]
                first = ws.max_column + 1
                for k, name in enumerate(cols):
                    c = ws.cell(row=head_row, column=first + k, value=name)
                    c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DCE8EA")
                    ws.column_dimensions[c.column_letter].width = 60 if k in (0, 4) else 28
                done_sheets[place["sheet"]] = first
            first = done_sheets[place["sheet"]]
            cid, name = parse_video(place["source"])
            video = videos.get(f"{cid} - {name}".strip(" -"))
            values = [res.get("answer", ""), res.get("sources", ""),
                      ("yes (declined)" if not res.get("answered") else "no (it answered)") if t.get("type") == "out_of_scope"
                      else facts.get(place["fact"], "-") if place["fact"] else "-",
                      ("yes" if video["found"] else "no") if video else "-", res.get("unsupported_claims", "") or "none",
                      fmt("faithfulness", res.get("faithfulness")), "PASS" if res.get("passed") else "CHECK", res.get("why", "")]
            for k, v in enumerate(values):
                c = ws.cell(row=place["row"], column=first + k, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if k == 6:
                    c.fill = PatternFill("solid", fgColor="E3F4E8" if v == "PASS" else "F7E3E1")
    wb.save(out)
    return out


def load_set(path: Path | None = None) -> list[dict]:
    path = path or qa_set_path()
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        if not path.exists():
            raise FileNotFoundError(f"No test plan at {path}")
        return load_test_plan(path)
    if not path.exists():
        raise FileNotFoundError(f"No test set at {path}. Make one: uv run qa_eval.py --draft 20 (then check it in Excel)")
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = [{k: (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]
    rows = [r for r in rows if r.get("question")]
    for i, r in enumerate(rows, 1):
        r["id"] = r.get("id") or f"q{i:02d}"
        r["type"] = r.get("type") if r.get("type") in TYPES else "specific"
    return rows


DRAFT = """You write test questions for a question-answering system over a company's training material. For each
excerpt write ONE question a learner could ask that this excerpt alone answers well, the expected answer (1-2
sentences, from the excerpt only) and 2-4 key facts (short phrases from the excerpt). Questions must be natural (as a
learner asks, not "according to the excerpt"), and about the content, not about the meeting. Skip an excerpt that is
small talk or logistics (return an empty question)."""

DRAFT_SCHEMA = {"type": "object", "properties": {"items": {"type": "array", "items": {"type": "object", "properties": {
    "excerpt": {"type": "string"}, "question": {"type": "string"}, "expected_answer": {"type": "string"},
    "key_facts": {"type": "array", "items": {"type": "string"}}}, "required": ["excerpt", "question"]}}},
    "required": ["items"]}

DRAFT_GENERAL = """You write test questions for a question-answering system over a company's training material. From
the list of sources (titles, descriptions, slide titles) write broad questions a learner could ask, each answered by
one source (e.g. "What is ...?", "How does ... work?"), with the source title. Then write questions in the same
field that these sources clearly do NOT answer (near the topic, but not covered), to test that the system declines."""

DRAFT_GENERAL_SCHEMA = {"type": "object", "properties": {
    "general": {"type": "array", "items": {"type": "object", "properties": {"question": {"type": "string"},
                                                                            "source": {"type": "string"}},
                                           "required": ["question", "source"]}},
    "out_of_scope": {"type": "array", "items": {"type": "string"}}}, "required": ["general", "out_of_scope"]}


def draft(n: int, path: Path | None = None, seed: int = 7) -> Path:
    """A first test set from the material: ~70% specific questions (each from one excerpt, so its
    source and time are known), ~20% general ones (per source) and ~10% (at least 3) that the
    material does not answer. Check and edit it before trusting the scores."""
    path = path or qa_set_path()
    every = list(store.all_segments().values())
    segs = [s for s in every if s.get("words", 0) >= 60] or [s for s in every if s.get("words", 0) >= 8]
    if not segs:
        raise ValueError("No processed material yet: run pipeline.py first")
    rng = random.Random(seed)
    by_source: dict[str, list[dict]] = {}
    for s in segs:
        by_source.setdefault(s["video_id"], []).append(s)
    n_specific = max(1, round(n * 0.7))
    picked: list[dict] = []
    pools = [rng.sample(v, len(v)) for v in by_source.values()]
    while len(picked) < n_specific and any(pools):   # spread over the sources
        for pool in pools:
            if pool and len(picked) < n_specific:
                picked.append(pool.pop())
    rows: list[dict] = []
    for i in range(0, len(picked), 5):
        batch = picked[i:i + 5]
        text = "\n\n".join(f"[E{k}] {s['title']} ({retrieval.where(s)}):\n{s['text']}" for k, s in enumerate(batch, 1))
        out = llm.chat_json(text, DRAFT_SCHEMA, system=DRAFT, step="qa draft", temperature=0.4)
        for item in out.get("items", []):
            m = re.fullmatch(r"E?(\d+)", str(item.get("excerpt", "")).strip())
            if not m or not item.get("question", "").strip() or not 1 <= int(m.group(1)) <= len(batch):
                continue
            s = batch[int(m.group(1)) - 1]
            at = "" if s.get("kind") == "document" else f"@{store.fmt_ts(s.get('start', 0))}"
            rows.append({"type": "specific", "question": item["question"].strip(),
                         "expected_answer": item.get("expected_answer", "").strip(),
                         "key_facts": "; ".join(x.strip() for x in item.get("key_facts", []) if x.strip()),
                         "expected_sources": f"{s['title']}{at}", "notes": f"drafted from {s['title']} {retrieval.where(s)}: check"})
    listing = []
    for items in by_source.values():
        first = items[0]
        slides = sorted({ln.strip() for s in items for ln in (s.get("slide_text") or "").splitlines()[:1] if ln.strip()})[:8]
        listing.append(f"- {first['title']}: {first.get('description', '')} Slides: {', '.join(slides)}")
    n_general = max(1, round(n * 0.2))
    n_out = max(3, n - n_specific - n_general)
    out = llm.chat_json(f"Write {n_general} general and {n_out} out-of-scope questions.\n\nSources:\n" + "\n".join(listing),
                        DRAFT_GENERAL_SCHEMA, system=DRAFT_GENERAL, step="qa draft", temperature=0.4)
    for g in out.get("general", [])[:n_general]:
        rows.append({"type": "general", "question": g["question"].strip(), "expected_answer": "", "key_facts": "",
                     "expected_sources": g.get("source", "").strip(), "notes": "drafted: add key facts"})
    for q in out.get("out_of_scope", [])[:n_out]:
        rows.append({"type": "out_of_scope", "question": str(q).strip(), "expected_answer": "", "key_facts": "",
                     "expected_sources": "", "notes": "drafted: check the material really doesn't cover it"})
    for i, r in enumerate(rows, 1):
        r.update(id=f"q{i:02d}", length="short")
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():   # never overwrite a set someone edited
        path = path.with_name(f"qa_set_draft_{datetime.now():%Y%m%d-%H%M%S}.csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SET_FIELDS)
        w.writeheader()
        w.writerows(rows)
    return path


# --- Run ---------------------------------------------------------------------------------------------------
def run(rows: list[dict], modes: list[str], model: str | None, judge: str, use_judge: bool = True,
        set_name: str = "qa_set.csv", plan: Path | None = None) -> Path:
    created = datetime.now().strftime("%Y%m%d-%H%M%S")
    folder = config.OUTPUT_DIR / f"{created}-qa-eval"
    folder.mkdir(parents=True, exist_ok=True)
    runs: dict[str, list[dict]] = {}
    for mode in modes:
        runs[mode] = []
        for k, row in enumerate(rows, 1):
            log.info("[%s %s/%s] %s", mode, k, len(rows), row["question"][:80])
            try:
                runs[mode].append(evaluate(row, mode, model, judge, use_judge))
            except Exception as e:   # one broken question must not stop the others: noted, and the run goes on
                log.error("%s failed: %s", row["id"], e, exc_info=not isinstance(e, (llm.LLMError, RuntimeError, OSError)))
                runs[mode].append({"id": row["id"], "type": row.get("type", ""), "question": row["question"], "mode": mode,
                                   "answered": False, "abstain_correct": False, "answer": f"(error: {e})", "sources": "",
                                   "invalid_labels": 0, "unsourced_sentences": 0, "numbers_not_in_sources": 0,
                                   "cost_usd": 0.0, "judge_cost_usd": 0.0, "seconds": 0.0, "passed": False,
                                   "why": f"error: {str(e)[:200]}", "folder": ""})
    with open(folder / "qa_results.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=RESULT_FIELDS, extrasaction="ignore")
        w.writeheader()
        for rows_ in runs.values():
            w.writerows([{**r, "missing_facts": "; ".join(r.get("missing_facts") or [])} for r in rows_])
    settings = {"created": created, "questions": len(rows), "set": set_name, "writer": model or config.CHAT_MODEL,
                "judge": judge, "rerank": config.RERANK_MODEL or config.CHAT_MODEL, "modes": modes}
    previous = None
    old = sorted(p for p in config.OUTPUT_DIR.glob("*-qa-eval") if p != folder and (p / "summary.json").exists())
    if old:
        previous = store.read_json(old[-1] / "summary.json", None)
    report(folder, runs, settings, previous)
    if plan and plan.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            write_back(plan, rows, runs[modes[-1]], folder / f"{plan.stem}_results.xlsx")
        except Exception as e:   # e.g. the plan is open in Excel: the report is still there
            log.error("Could not write the results into a copy of %s (%s); see qa_report.html and qa_results.csv",
                      plan.name, e)
    store.write_json(folder / "summary.json", {**summary(runs[modes[-1]]), "created": created, "mode": modes[-1],
                                               "by_mode": {m: summary(r) for m, r in runs.items()}, "settings": settings})
    return folder


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--draft", type=int, metavar="N", help="write a first test set of about N questions from the material")
    ap.add_argument("--set", type=Path, help="the test set (default eval/qa_set.csv)")
    ap.add_argument("--compare", action="store_true", help="every retrieval mode: meaning, hybrid, hybrid+rerank")
    ap.add_argument("--mode", choices=retrieval.MODES, help="one retrieval mode (default RETRIEVAL in .env)")
    ap.add_argument("--limit", type=int, help="only the first N questions")
    ap.add_argument("--ids", help="only these ids, comma-separated")
    ap.add_argument("--model", help="the writer (default CHAT_MODEL)")
    ap.add_argument("--judge", help="the judge (default JUDGE_MODEL, else REVIEW_MODEL)")
    ap.add_argument("--no-judge", action="store_true", help="only the code checks (no model grading; cheaper)")
    args = ap.parse_args()
    config.setup_logging()
    args.set = args.set or qa_set_path()
    try:
        if args.draft:
            path = draft(args.draft, args.set)
            print(f"Draft written: {path}\nCheck it in Excel (questions, key facts, expected sources), add your own, "
                  f"then run: uv run qa_eval.py" + (f" --set \"{path}\"" if path != args.set else ""))
            return 0
        rows = load_set(args.set)
        if args.ids:
            wanted = {x.strip() for x in args.ids.split(",")}
            rows = [r for r in rows if r["id"] in wanted]
        if args.limit:
            rows = rows[:args.limit]
        modes = list(retrieval.MODES) if args.compare else [args.mode or config.RETRIEVAL]
        folder = run(rows, modes, args.model, judge_model(args.judge), not args.no_judge, args.set.name, args.set)
    except (FileNotFoundError, ValueError, llm.LLMError) as e:
        log.error("%s", e)
        return 1
    s = store.read_json(folder / "summary.json", {})
    for m, x in s.get("by_mode", {}).items():
        print(f"{m:<14} passed {x['passed']}/{x['questions']}  faithfulness {fmt('faithfulness', x['faithfulness'])}  "
              f"hallucination {fmt('hallucination_rate', x['hallucination_rate'])}  retrieval hit "
              f"{fmt('retrieval_hit', x['retrieval_hit'])}  key facts {fmt('key_fact_recall', x['key_fact_recall'])}  "
              f"declined out-of-scope {fmt('declined_out_of_scope', x['declined_out_of_scope'])}")
    print(f"\nReport: {folder / 'qa_report.html'}\nRate in Excel: {folder / 'qa_results.csv'}")
    for x in folder.glob("*_results.xlsx"):
        print(f"Your test plan with the results: {x}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
